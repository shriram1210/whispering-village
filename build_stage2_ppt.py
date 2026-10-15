"""build_stage2_ppt.py - Stage 2 build deck: what we built, uniqueness, real-world problem."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

HERE = os.path.dirname(os.path.abspath(__file__))
BG = RGBColor(20, 22, 36)
GOLD = RGBColor(255, 220, 110)
WHITE = RGBColor(240, 240, 245)
GREY = RGBColor(150, 150, 170)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    tx = s.shapes.add_textbox(Inches(0.4), Inches(7.05), Inches(12.5), Inches(0.35))
    tx.text_frame.text = "Whispering Village  •  Stage 2 build  •  Cognitive NPCs"
    for p in tx.text_frame.paragraphs:
        p.alignment = PP_ALIGN.RIGHT
        for r in p.runs:
            r.font.size = Pt(10)
            r.font.color.rgb = GREY
    return s


def title(s, text, size=34):
    tx = s.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(12.1), Inches(1.0))
    p = tx.text_frame.paragraphs[0]
    p.text = text
    r = p.runs[0]
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = GOLD


def bullets(s, items, width=11.9, size=20):
    tx = s.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(width), Inches(5.0))
    tf = tx.text_frame
    tf.word_wrap = True
    for i, line in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        r = p.add_run()
        r.text = "•  " + line
        r.font.size = Pt(size)
        r.font.color.rgb = WHITE


def shot(s, path, w=5.6):
    full = os.path.join(HERE, path)
    if os.path.exists(full):
        s.shapes.add_picture(full, Inches(13.333 - w - 0.5), Inches(1.5), width=Inches(w))


# 1. Title
s = slide()
t = s.shapes.add_textbox(Inches(0.6), Inches(1.8), Inches(12.1), Inches(1.2))
p = t.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Whispering Village — Stage 2"
r.font.size = Pt(56)
r.font.bold = True
r.font.color.rgb = GOLD
t2 = s.shapes.add_textbox(Inches(0.6), Inches(3.2), Inches(12.1), Inches(1.6))
p = t2.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "A playable village where conversations change the game\nTalk • earn trust • sleep • watch rumors spread"
r.font.size = Pt(24)
r.font.color.rgb = WHITE

# 2. What we built
s = slide()
title(s, "Stage 2 — what we built and verified")
bullets(s, [
    "Playable twilight village: walk, talk to Gruff / Mira / Rex, quest, shop, gate",
    "Trust engine: one score drives quests, prices, secrets and the gate",
    "SQLite memory: facts + nightly summaries, survives restarts (19-check suite green)",
    "Voice mode: mic button + V key, tone shifts trust, Hinglish understood",
    "Stage 3 gossip engine + rumor map already live on top of it",
    "Every rule unit-tested with stubbed AI — zero quota burned in tests",
], width=6.6)
shot(s, "demo/walk.png", w=5.9)

# 3. How it works
s = slide()
title(s, "How it works — AI suggests, game decides")
bullets(s, [
    "Player types/speaks → LLM replies in character + proposes an action (JSON)",
    "world.py validates against trust + quest state — the AI can never invent rewards",
    "Night: gossip.py retells memories villager-to-villager, distorted by personality",
    "Offline fallback at every layer — the full loop demos with no internet",
    "Keys live only in .env; check_api.py diagnoses the AI link in seconds",
])

# 4. Uniqueness
s = slide()
title(s, "What makes it unique")
bullets(s, [
    "Mutating gossip: Mira doubles numbers, Gruff grunts, Rex files reports — lies evolve",
    "Rumor map: live graph of who-believes-what, click any villager to inspect",
    "Tone-aware NPCs: loud/fast reads rude (−6), quiet/slow reads sad (+3 sympathy)",
    "Forgetting-proof memory: 10 or 10,000 chats cost the same prompt",
    "Trust economy: kindness is literally currency (prices, quests, gates, secrets)",
], width=6.6)
shot(s, "demo/gossip.png", w=5.9)

# 5. Real-world problem
s = slide()
title(s, "Real-world problem it solves")
bullets(s, [
    "Misinformation: rumors outrun facts — our village is a playable inoculation lab (Bad News game precedent, N=15k)",
    "Language access: code-switched players get spoken-back Hinglish, not broken English (Duolingo Max precedent)",
    "Empathy practice: kindness-gated quests reward how you treat people, safely rehearsable",
    "Low-literacy play: voice-first input, icon buttons, short sentences",
])

# 6. How Stage 2 was uniquely built
s = slide()
title(s, "How we built Stage 2 differently")
bullets(s, [
    "Deterministic core, AI at the edges: rules tested headless, LLM only flavors",
    "Unity-to-Pygame art pipeline: characters batch-rendered in Unity 6, chroma-keyed in (SpriteBaker.cs)",
    "Mic you can see: procedural mic button + REC pulse + device diagnostics, no assets",
    "Proof over promises: GOSSIP_OK, BACKEND_OK, THREAD_OK suites + rendered screenshots from real code",
], width=6.6)
shot(s, "demo/chat.png", w=5.9)

# 7. Demo + next
s = slide()
title(s, "Live demo + what is next")
bullets(s, [
    "Demo: greet Mira → earn Gruff's quest → open Rex's gate → fetch ore → Sleep → open the rumor map",
    "Contrast: kind voice vs arrogant voice — same words, different village",
    "Proof: /status diary, restart persistence, offline mode with wifi killed",
    "Next: pixel font, Vosk offline speech-to-text, 2 more villagers, rumor-map overlay polish",
])

out = os.path.join(HERE, "Whispering_Village_Stage2.pptx")
prs.save(out)
print("saved:", out, "| slides:", len(list(prs.slides)))
