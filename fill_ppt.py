"""fill_ppt.py - fills round 1 template with Whispering Village content. Theme untouched."""
import os
from pptx import Presentation

SRC = r"C:\Users\shrir\Downloads\round 1 template.pptx"
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Whispering_Village_Round1.pptx")

C = {
    (0, "TextBox 21"): ["TEAM: Team Whispering Village     •     MISSION: 05 — Cognitive NPCs"],
    (1, "TextBox 17"): ["• NPCs still use fixed dialogue trees and forget everything",
                        "• Mission 5 asks for cognitive NPCs: natural talk that drives quests",
                        "• ...and remembers the player across sessions"],
    (1, "TextBox 21"): ["• Players lose immersion; learners get no safe practice space",
                        "• Real-world link: misinformation defense, language access, empathy training",
                        "• A village that teaches whom to trust helps all three"],
    (1, "TextBox 25"): ["• Scope: 1 village, 3 NPCs, text + voice, SQLite memory",
                        "• Assume: laptop hardware, free-tier LLM + offline fallback",
                        "• Validate: trust-gated quest, memory after restart"],
    (2, "TextBox 17"): ["• Whispering Village: Gruff (smith), Mira (merchant), Rex (guard)",
                        "• They converse naturally, score trust, gossip overnight",
                        "• Quests, prices and gates change by what — and how — you speak"],
    (2, "TextBox 21"): ["• Type/speak → LLM replies in character + proposes an action",
                        "• Game validates against trust rules; memory is saved",
                        "• Nights: NPCs summarize the day; voice tone shifts trust"],
    (2, "TextBox 25"): ["• Walkable village + 3 talking NPCs + 1 dialogue-driven quest",
                        "• Restart-proof memory + 3-minute judge demo",
                        "• Whole loop demoable fully offline"],
    (3, "TextBox 17"): ["• Typed text, mic voice (V key / mic button)",
                        "• Player position + proximity to NPCs",
                        "• Trust, quest, inventory and memory state"],
    (3, "TextBox 21"): ["• npc.py prompts (persona + memories + game context) → LLM JSON",
                        "• voice.py prosody → expression → trust delta",
                        "• world.py validates every AI-proposed action (AI suggests, game decides)"],
    (3, "TextBox 25"): ["• In-character replies with mood, trust bar, portraits",
                        "• Quest / price / gate changes on screen",
                        "• Rumor-verification quests, diary summaries"],
    (4, "TextBox 17"): ["• None required — pure software build",
                        "• Dev + player laptop only",
                        "• Optional microphone for voice mode"],
    (4, "TextBox 21"): ["• Python 3.14, Pygame-ce, SQLite FTS5",
                        "• google-genai (Gemini) + OpenAI-compatible OpenCode gateway",
                        "• SpeechRecognition, sounddevice, VADER sentiment"],
    (4, "TextBox 25"): ["• data/npcs.json (personalities), game_memory.db (runtime memory)",
                        "• Free-tier LLM + built-in offline replies — no datasets needed",
                        "• check_api.py diagnoses the AI link in seconds"],
    (5, "TextBox 17"): ["• Free-tier quota holds for a 3-minute demo",
                        "• pygame-ce wheels exist for Python 3.14 (verified)",
                        "• 2D scope stays fixed; no 3D engine risk"],
    (5, "TextBox 21"): ["• LLM latency / model retirements → multi-model fallback + timeout",
                        "• No mic on demo laptop → typed mode covers everything",
                        "• Bad AI output → game validates; AI can never invent rewards"],
    (5, "TextBox 25"): ["• Fully offline mode demos the whole loop (built-in replies)",
                        "• Terminal-chat fallback if Pygame fails",
                        "• Pre-recorded demo video as last resort"],
    (6, "TextBox 17"): ["• Milestone: village scene, movement, dialogue panel (M1+M2)",
                        "• Target: walk + talk with placeholder replies"],
    (6, "TextBox 21"): ["• Milestone: real AI replies, trust system, quest/gate/shop rules (M2+M3)",
                        "• Target: quest completable purely by talking"],
    (6, "TextBox 25"): ["• Milestone: SQLite memory + summaries, voice mode, rumor quest (all)",
                        "• Target: restart-proof memory + tone-based trust"],
    (6, "TextBox 29"): ["• Milestone: headless tests, quota-safe demo script, docs (M4+all)",
                        "• Target: reliable 3-minute demo + this deck"],
    (6, "TextBox 33"): ["• Member 1: Pygame world + art   • Member 2: AI dialogue + voice",
                        "• Member 3: memory + quests + rules   • Member 4: integration, tests, demo"],
    (7, "TextBox 17"): ["• Free-text talk, no fixed options",
                        "• Quest completes only via validated dialogue",
                        "• Memory provable after restart; voice visibly moves trust; runs offline"],
    (7, "TextBox 21"): ["• 1. Name → quest → Rex gate → glowing ore → sword",
                        "• 2. Kind vs arrogant voice contrast on trust",
                        "• 3. /sleep summary + rumor-verification quest"],
    (7, "TextBox 25"): ["• How is uniqueness weighted vs polish in judging?",
                        "• Is a fully-offline demo acceptable if venue net fails?",
                        "• Advice on scaling gossip from 3 to 5 NPCs?"],
}


def set_box(slide, name, lines):
    for sh in slide.shapes:
        if sh.name == name and sh.has_text_frame:
            tf = sh.text_frame
            tf.clear()
            for i, line in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = line
                p.level = 0
            return True
    return False


prs = Presentation(SRC)
missed = []
for (si, box), lines in C.items():
    if not set_box(prs.slides[si], box, lines):
        missed.append((si, box))
prs.save(OUT)
print("saved:", OUT, "| missed:", missed or "none")
